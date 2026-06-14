#!/usr/bin/env python3
"""
docs/slides/build_deck.py — Generate the Bespoke Agentic AI Factory demo deck.

Renders 35 slides (title, architecture, flows, UI mockups) as high-resolution
PNGs with matplotlib, then assembles them into:
  - docs/slides/ai-factory-demo.pdf
  - docs/slides/ai-factory-demo.pptx   (one full-bleed image per slide)

Fully offline — no browser / LibreOffice required. UI "screenshots" are
high-fidelity mockups (the live stack needs Docker).

Run:  python docs/slides/build_deck.py
"""
from __future__ import annotations

import pathlib

import matplotlib
matplotlib.use("Agg")
import matplotlib.font_manager as fm  # noqa: E402
import numpy as np  # noqa: E402
from matplotlib import pyplot as plt  # noqa: E402
from matplotlib.patches import Circle, FancyBboxPatch, Rectangle  # noqa: E402

HERE = pathlib.Path(__file__).resolve().parent
ASSETS = HERE / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)

# ── Canvas & palette ─────────────────────────────────────────────────────────
W, H = 1280, 720

BG   = "#0B1020"   # deep navy
BG2  = "#111A2E"   # panel
BG3  = "#0E1626"
INK  = "#E9F0FB"   # near-white
MUT  = "#93A4C2"   # muted blue-gray
LINE = "#27344F"
CYAN = "#22D3EE"
BLUE = "#3B82F6"
PUR  = "#A78BFA"
GRN  = "#34D399"
AMB  = "#FBBF24"
RED  = "#F87171"
ORG  = "#FF9900"   # AWS
TEAL = "#2DD4BF"
PINK = "#F472B6"

STAGE = [("Retrieve", CYAN), ("Enrich", PUR), ("Transform", ORG),
         ("Generate", BLUE), ("Validate", GRN)]

try:
    MONO = fm.FontProperties(family="DejaVu Sans Mono")
except Exception:                                       # pragma: no cover
    MONO = fm.FontProperties(family="monospace")

_slides: list[pathlib.Path] = []


# ── Primitives ───────────────────────────────────────────────────────────────

def new_slide(bg=BG):
    fig = plt.figure(figsize=(12.8, 7.2), dpi=150)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W)
    ax.set_ylim(H, 0)               # flip: (0,0) top-left, screen coords
    ax.axis("off")
    ax.add_patch(Rectangle((0, 0), W, H, color=bg, zorder=-20))
    return fig, ax


def grad_band(ax, y=0, h=8, c=CYAN):
    ax.add_patch(Rectangle((0, y), W, h, color=c, zorder=5))


def footer(ax, n, total=35):
    ax.plot([56, W - 56], [H - 46, H - 46], color=LINE, lw=1, zorder=2)
    ax.text(56, H - 28, "Bespoke Agentic AI Factory  ·  Demo Pack",
            color=MUT, fontsize=9, va="center")
    ax.text(W - 56, H - 28, f"{n:02d} / {total:02d}",
            color=MUT, fontsize=9, ha="right", va="center")


def title_bar(ax, kicker, title, sub=None):
    ax.add_patch(Rectangle((56, 52), 46, 6, color=CYAN, zorder=3))
    ax.text(112, 56, kicker.upper(), color=CYAN, fontsize=12.5,
            fontweight="bold", va="center")
    ax.text(56, 96, title, color=INK, fontsize=29, fontweight="bold", va="center")
    if sub:
        ax.text(56, 134, sub, color=MUT, fontsize=14, va="center")


def rbox(ax, x, y, w, h, fc, ec=None, lw=1.6, r=14, alpha=1.0, z=3):
    ax.add_patch(FancyBboxPatch(
        (x, y), w, h, boxstyle=f"round,pad=0,rounding_size={r}",
        fc=fc, ec=ec or fc, lw=lw, alpha=alpha, zorder=z, mutation_aspect=1))


def boxed_label(ax, x, y, w, h, title, fc, ec, tc=INK, sub=None,
                title_fs=15, sub_fs=11, z=4):
    rbox(ax, x, y, w, h, fc=fc, ec=ec, lw=2, z=z)
    cy = y + h / 2 if not sub else y + h / 2 - 10
    ax.text(x + w / 2, cy, title, color=tc, fontsize=title_fs,
            fontweight="bold", ha="center", va="center", zorder=z + 1)
    if sub:
        ax.text(x + w / 2, y + h / 2 + 14, sub, color=tc, fontsize=sub_fs,
                ha="center", va="center", alpha=0.85, zorder=z + 1)


def chip(ax, x, y, text, fc, tc=INK, fs=11, padx=14, h=26, z=6):
    w = padx * 2 + len(text) * fs * 0.74
    rbox(ax, x, y, w, h, fc=fc, ec=fc, lw=0, r=h / 2, z=z)
    ax.text(x + w / 2, y + h / 2, text, color=tc, fontsize=fs,
            fontweight="bold", ha="center", va="center", zorder=z + 1)
    return w


def arrow(ax, x1, y1, x2, y2, color=MUT, lw=2.2, style="-|>", z=2, ls="-"):
    ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                arrowprops=dict(arrowstyle=style, color=color, lw=lw,
                                shrinkA=3, shrinkB=3, linestyle=ls), zorder=z)


def bullets(ax, x, y, items, dy=52, fs=15, color=INK, marker=CYAN, mw=560):
    for i, it in enumerate(items):
        yy = y + i * dy
        ax.add_patch(FancyBboxPatch((x, yy - 6), 12, 12,
                     boxstyle="round,pad=0,rounding_size=3",
                     fc=marker, ec=marker, zorder=4))
        ax.text(x + 26, yy, it, color=color, fontsize=fs, va="center", zorder=4,
                wrap=True)


def panel(ax, x, y, w, h, fc=BG2, ec=LINE, r=16, lw=1.4, z=2):
    rbox(ax, x, y, w, h, fc=fc, ec=ec, lw=lw, r=r, z=z)


def mock_window(ax, x, y, w, h, url=None, title=None, body=BG3):
    rbox(ax, x, y, w, h, fc=body, ec=LINE, lw=1.6, r=12, z=3)
    bar_h = 36
    ax.add_patch(FancyBboxPatch((x, y), w, bar_h + 14,
                 boxstyle="round,pad=0,rounding_size=12",
                 fc="#15203A", ec=LINE, lw=1.4, zorder=4))
    ax.add_patch(Rectangle((x, y + bar_h - 4), w, 6, color="#15203A", zorder=4))
    for i, c in enumerate(("#FF5F57", "#FEBC2E", "#28C840")):
        ax.add_patch(Circle((x + 22 + i * 22, y + bar_h / 2), 6.5, color=c, zorder=6))
    if url is not None:
        rbox(ax, x + 96, y + 8, w - 130, bar_h - 16, fc="#0B1426", ec=LINE, lw=1, r=8, z=5)
        ax.text(x + 112, y + bar_h / 2, url, color=MUT, fontsize=10,
                va="center", zorder=6, fontproperties=MONO)
    if title is not None:
        ax.text(x + w / 2, y + bar_h / 2, title, color=INK, fontsize=11,
                fontweight="bold", ha="center", va="center", zorder=6)
    return x + 14, y + bar_h + 14, w - 28, h - bar_h - 28


def mini_chart(ax, x, y, w, h, series, colors, title, fill=False, unit=""):
    panel(ax, x, y, w, h, fc="#0C1424", ec=LINE, r=10, z=5)
    ax.text(x + 14, y + 18, title, color=INK, fontsize=11, fontweight="bold",
            va="center", zorder=7)
    px, py = x + 14, y + 34
    pw, ph = w - 28, h - 48
    for k in range(1, 4):                       # gridlines
        gy = py + ph * k / 4
        ax.plot([px, px + pw], [gy, gy], color=LINE, lw=0.6, zorder=6)
    for s, c in zip(series, colors):
        t = np.linspace(0, 1, len(s))
        X = px + t * pw
        Y = (py + ph) - np.array(s) * ph
        ax.plot(X, Y, color=c, lw=2.2, zorder=8)
        if fill:
            ax.fill_between(X, Y, py + ph, color=c, alpha=0.12, zorder=7)


def stat_tile(ax, x, y, w, h, value, label, color, good=True):
    panel(ax, x, y, w, h, fc="#0C1424", ec=LINE, r=10, z=5)
    ax.add_patch(Rectangle((x, y + 10), 5, h - 20, color=color, zorder=6))
    ax.text(x + 18, y + h / 2 - 6, value, color=color, fontsize=23,
            fontweight="bold", va="center", zorder=7)
    ax.text(x + 18, y + h - 18, label, color=MUT, fontsize=9.5, va="center", zorder=7)


def terminal(ax, x, y, w, h, lines, title="bash — ai-factory"):
    ix, iy, iw, ih = mock_window(ax, x, y, w, h, title=title, body="#05080F")
    ln = iy + 6
    for kind, text in lines:
        col = {"p": GRN, "c": INK, "o": MUT, "ok": GRN, "warn": AMB,
               "hdr": CYAN, "err": RED}.get(kind, INK)
        prefix = "$ " if kind == "c" else ""
        s = ("➜ " if kind == "c" else "") + text
        ax.text(ix + 8, ln, s, color=col, fontsize=10.5, va="center",
                zorder=8, fontproperties=MONO)
        ln += 22
    return ix, iy, iw, ih


def save(fig, n):
    path = ASSETS / f"slide_{n:02d}.png"
    fig.savefig(path, dpi=150, facecolor=BG)
    plt.close(fig)
    _slides.append(path)


# ════════════════════════════════════════════════════════════════════════════
# SLIDES
# ════════════════════════════════════════════════════════════════════════════

def s01():
    fig, ax = new_slide()
    # decorative pipeline ribbon
    for i, (_, c) in enumerate(STAGE):
        ax.add_patch(Rectangle((0, 150 + i * 84), W, 84, color=c, alpha=0.05, zorder=-5))
    grad_band(ax, 0, 8, CYAN)
    ax.text(64, 250, "Bespoke", color=INK, fontsize=58, fontweight="bold", va="center")
    ax.text(64, 320, "Agentic AI Factory", color=CYAN, fontsize=58,
            fontweight="bold", va="center")
    ax.text(66, 380, "Composable · Observable · Governed AI pipelines for the enterprise",
            color=MUT, fontsize=18, va="center")
    cx = 66
    for t, c in [("Agentic RAG", CYAN), ("AWS Kiro", PUR), ("AWS Transform", ORG),
                 ("Spring AI", GRN), ("DeepEval Gates", AMB)]:
        cx += chip(ax, cx, 430, t, fc="#152038", tc=c, fs=12) + 12
    ax.text(64, 540, "Solution Architecture · Pipeline Flows · Live Demo",
            color=INK, fontsize=15, va="center")
    ax.text(64, 568, "35-slide demo pack", color=MUT, fontsize=12, va="center")
    footer(ax, 1)
    save(fig, 1)


def s02():
    fig, ax = new_slide()
    title_bar(ax, "Executive summary", "One factory, five governed stages")
    bullets(ax, 64, 210, [
        "Turn scattered enterprise knowledge into grounded, cited answers",
        "Every answer passes automated quality + safety gates before release",
        "Modernise legacy Java/Spring estates with agentic transforms",
        "Full observability: metrics, traces and a single pane of glass",
        "Ships on AWS (ECS Fargate) via Terraform + Harness with canary releases",
    ], dy=58, fs=16)
    # right metric cards
    cards = [("5", "governed pipeline stages", CYAN),
             ("2", "internal sources wired (Confluence, SharePoint)", PUR),
             ("99", "automated tests, all green", GRN),
             ("4", "deploy environments w/ canary", ORG)]
    bx = 760
    for i, (v, l, c) in enumerate(cards):
        y = 200 + i * 96
        panel(ax, bx, y, 440, 80, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((bx, y + 10), 5, 60, color=c, zorder=4))
        ax.text(bx + 26, y + 40, v, color=c, fontsize=30, fontweight="bold", va="center")
        ax.text(bx + 110, y + 40, l, color=INK, fontsize=13, va="center")
    footer(ax, 2)
    save(fig, 2)


def s03():
    fig, ax = new_slide()
    title_bar(ax, "The problem", "Enterprise AI stalls on trust, silos and legacy")
    cols = [
        ("Knowledge silos", RED,
         ["Docs spread across Confluence,", "SharePoint, wikis and drives",
          "No unified, searchable context"]),
        ("Unverified output", AMB,
         ["LLMs hallucinate and drift", "No automated quality gate",
          "Compliance / audit risk"]),
        ("Legacy estates", ORG,
         ["Old Spring Boot, CVE exposure", "Hardcoded rule engines",
          "Manual, slow modernisation"]),
    ]
    for i, (t, c, items) in enumerate(cols):
        x = 64 + i * 392
        panel(ax, x, 200, 360, 360, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((x, 200), 360, 8, color=c, zorder=4))
        ax.text(x + 24, 250, t, color=INK, fontsize=19, fontweight="bold", va="center")
        for j, it in enumerate(items):
            ax.add_patch(Circle((x + 30, 308 + j * 50), 4, color=c, zorder=5))
            ax.text(x + 48, 308 + j * 50, it, color=MUT, fontsize=13.5, va="center")
    ax.text(64, 600, "→  The gap: a governed, observable pipeline that grounds, "
            "verifies and modernises — not just a chatbot.",
            color=CYAN, fontsize=15, va="center", fontweight="bold")
    footer(ax, 3)
    save(fig, 3)


def s04():
    fig, ax = new_slide()
    title_bar(ax, "Solution", "Capabilities at a glance")
    caps = [
        ("Agentic RAG", CYAN, "Hybrid vector + graph,\nmulti-hop retrieval"),
        ("Source connectors", PUR, "Confluence & SharePoint\n→ one RAG index"),
        ("Eval gates", GRN, "DeepEval metrics +\nLLM-as-judge, CI-gated"),
        ("AWS Transform", ORG, "7 modernisation\ndefinitions (atx)"),
        ("Spring AI", BLUE, "Boot 2.x → 3 +\nSpring AI 1.0"),
        ("Single pane of glass", TEAL, "Metrics, traces,\nRED per stage"),
        ("Secure delivery", PINK, "SAST/DAST gates,\ncanary releases"),
        ("Knowledge graph", AMB, "Neo4j entity\nenrichment"),
    ]
    for i, (t, c, d) in enumerate(caps):
        r, col = divmod(i, 4)
        x = 64 + col * 296
        y = 200 + r * 190
        panel(ax, x, y, 272, 166, fc=BG2, ec=LINE)
        ax.add_patch(FancyBboxPatch((x + 20, y + 20), 44, 44,
                     boxstyle="round,pad=0,rounding_size=10", fc="#152038", ec=c, lw=2, zorder=4))
        ax.add_patch(Circle((x + 42, y + 42), 9, color=c, zorder=5))
        ax.text(x + 20, y + 92, t, color=INK, fontsize=15, fontweight="bold", va="center")
        ax.text(x + 20, y + 122, d, color=MUT, fontsize=11.5, va="top")
    footer(ax, 4)
    save(fig, 4)


def s05():
    fig, ax = new_slide()
    title_bar(ax, "Architecture", "High-level system architecture")
    # client
    boxed_label(ax, 56, 210, 150, 70, "Client", BG2, BLUE, sub="Spring API :8080")
    arrow(ax, 206, 245, 250, 245, color=MUT)
    # orchestrator
    boxed_label(ax, 250, 200, 150, 90, "Kiro Agent", BG2, PUR, sub="pipeline spec\n+ retries")
    # stage column
    sx = 452
    for i, (name, c) in enumerate(STAGE):
        y = 168 + i * 78
        boxed_label(ax, sx, y, 196, 60, f"{i+1}. {name}", "#0E1A30", c, tc=c, title_fs=14)
        if i:
            arrow(ax, sx + 98, 168 + (i - 1) * 78 + 60, sx + 98, y, color=LINE, lw=1.6)
    arrow(ax, 400, 245, sx, 245, color=MUT)
    # right: datastores + model + eval
    rx = 700
    boxed_label(ax, rx, 176, 200, 56, "OpenSearch", BG2, CYAN, sub="hybrid vectors")
    boxed_label(ax, rx, 246, 200, 56, "Neo4j", BG2, PUR, sub="knowledge graph")
    boxed_label(ax, rx, 316, 200, 56, "AWS Transform", BG2, ORG, sub="atx custom def")
    boxed_label(ax, rx, 386, 200, 56, "Claude (Sonnet)", BG2, BLUE, sub="generation")
    boxed_label(ax, rx, 456, 200, 56, "Eval Service", BG2, GRN, sub="DeepEval + judge")
    for i, yy in enumerate((204, 274, 344, 414, 484)):
        arrow(ax, sx + 196, 168 + i * 78 + 30, rx, yy, color=LINE, lw=1.4)
    # connectors feeding retrieve
    boxed_label(ax, rx + 220, 176, 180, 56, "Confluence", BG2, PINK, sub="REST connector")
    boxed_label(ax, rx + 220, 246, 180, 56, "SharePoint", BG2, TEAL, sub="Graph connector")
    arrow(ax, rx + 220, 204, rx + 200, 204, color=PINK, lw=1.6, style="<|-")
    arrow(ax, rx + 220, 274, rx + 200, 274, color=TEAL, lw=1.6, style="<|-")
    # observability strip
    panel(ax, 56, 556, 1064, 96, fc=BG3, ec=LINE)
    ax.text(76, 580, "OBSERVABILITY", color=CYAN, fontsize=11, fontweight="bold", va="center")
    obs = [("OTel Collector", CYAN), ("Prometheus", AMB), ("Jaeger", PUR),
           ("Grafana", ORG), ("CloudWatch", BLUE)]
    ox = 76
    for t, c in obs:
        ox += chip(ax, ox, 600, t, fc="#152038", tc=c, fs=12) + 14
    ax.text(W - 76, 600, "AWS · ECS Fargate · ALB · Terraform · Harness",
            color=MUT, fontsize=12, ha="right", va="center", fontproperties=MONO)
    footer(ax, 5)
    save(fig, 5)


def s06():
    fig, ax = new_slide()
    title_bar(ax, "Core flow", "The 5-stage agentic pipeline")
    y = 250
    bw, gap = 196, 36
    x = 64
    notes = ["hybrid search\nOpenSearch", "graph context\nNeo4j",
             "normalise + PII\nAWS Transform", "grounded answer\nClaude", "metrics + judge\nDeepEval"]
    for i, (name, c) in enumerate(STAGE):
        boxed_label(ax, x, y, bw, 92, f"{i+1}. {name}", "#0E1A30", c, tc=c, title_fs=17)
        ax.text(x + bw / 2, y + 124, notes[i], color=MUT, fontsize=11.5, ha="center", va="top")
        if i < 4:
            arrow(ax, x + bw, y + 46, x + bw + gap, y + 46, color=c, lw=2.6)
        x += bw + gap
    # loop-back for multi-hop
    arrow(ax, 64 + bw / 2, y, 64 + bw / 2, y - 44, color=CYAN, lw=2)
    ax.annotate("", xy=(64 + bw / 2, y), xytext=(64 + 1.5 * bw + gap, y - 44),
                arrowprops=dict(arrowstyle="-|>", color=CYAN, lw=2,
                                connectionstyle="arc3,rad=0.3"), zorder=2)
    ax.text(64 + bw, y - 56, "multi-hop: re-retrieve until sufficiency met",
            color=CYAN, fontsize=12, ha="center", va="center", fontweight="bold")
    panel(ax, 64, 470, 1056, 150, fc=BG2, ec=LINE)
    ax.text(86, 500, "Why it matters", color=INK, fontsize=15, fontweight="bold", va="center")
    bullets(ax, 86, 540, [
        "Each stage is an independent, instrumented service (OTel span per stage)",
        "Retry + sufficiency scoring make retrieval agentic, not single-shot",
        "The validate stage can block a bad answer before it reaches the user",
    ], dy=34, fs=13.5)
    footer(ax, 6)
    save(fig, 6)


def s07():
    fig, ax = new_slide()
    title_bar(ax, "Deep dive", "Agentic RAG — hybrid retrieval + graph")
    # left flow
    boxed_label(ax, 64, 220, 230, 64, "Query", BG2, INK)
    boxed_label(ax, 64, 320, 230, 64, "Dense kNN (BGE)", "#0E1A30", CYAN, tc=CYAN)
    boxed_label(ax, 64, 410, 230, 64, "BM25 sparse", "#0E1A30", AMB, tc=AMB)
    arrow(ax, 179, 284, 179, 320, color=MUT)
    arrow(ax, 179, 284, 179, 410, color=MUT, ls=":")
    boxed_label(ax, 340, 365, 180, 70, "RRF fusion", "#0E1A30", PUR, tc=PUR, sub="k = 60")
    arrow(ax, 294, 352, 340, 392, color=CYAN)
    arrow(ax, 294, 442, 340, 410, color=AMB)
    boxed_label(ax, 560, 365, 200, 70, "Neo4j enrich", "#0E1A30", GRN, tc=GRN, sub="related entities")
    arrow(ax, 520, 400, 560, 400, color=PUR)
    boxed_label(ax, 800, 320, 200, 70, "Sufficiency?", BG2, CYAN, tc=CYAN, sub="Claude check")
    arrow(ax, 760, 400, 800, 360, color=GRN)
    boxed_label(ax, 800, 430, 200, 64, "Answer + sources", "#0E1A30", BLUE, tc=BLUE)
    arrow(ax, 900, 390, 900, 430, color=GRN)
    arrow(ax, 800, 340, 179, 256, color=CYAN, lw=1.6,
          style="-|>",)
    ax.text(500, 250, "loop: re-retrieve if score below threshold",
            color=CYAN, fontsize=11.5, va="center")
    panel(ax, 64, 540, 1056, 96, fc=BG2, ec=LINE)
    for i, (t, d, c) in enumerate([
        ("Hybrid", "dense + sparse, best of both", CYAN),
        ("Grounded", "every claim cites a source", GRN),
        ("Agentic", "multi-hop with sufficiency", PUR),
        ("Filtered", "metadata + dept scoping", AMB)]):
        x = 86 + i * 262
        ax.text(x, 572, t, color=c, fontsize=15, fontweight="bold", va="center")
        ax.text(x, 602, d, color=MUT, fontsize=12, va="center")
    footer(ax, 7)
    save(fig, 7)


def s08():
    fig, ax = new_slide()
    title_bar(ax, "New capability", "Internal source connectors",
              "Confluence & SharePoint → the agentic RAG index")
    # sources
    boxed_label(ax, 64, 250, 200, 80, "Confluence", BG2, PINK, sub="Atlassian REST")
    boxed_label(ax, 64, 360, 200, 80, "SharePoint", BG2, TEAL, sub="MS Graph")
    boxed_label(ax, 64, 470, 200, 64, "+ pluggable", BG3, MUT, tc=MUT, sub="implement Connector")
    # framework
    fx = 330
    panel(ax, fx, 230, 320, 320, fc=BG2, ec=LINE)
    ax.text(fx + 20, 262, "Connector framework", color=INK, fontsize=15, fontweight="bold", va="center")
    for i, (t, c) in enumerate([("fetch()", CYAN), ("html → text", AMB),
                                ("chunk()", PUR), ("SourceDocument", GRN),
                                ("offline fallback (sample data)", ORG)]):
        y = 300 + i * 48
        ax.add_patch(Circle((fx + 30, y), 5, color=c, zorder=5))
        ax.text(fx + 48, y, t, color=INK, fontsize=13.5, va="center", fontproperties=MONO)
    for sy in (290, 400):
        arrow(ax, 264, sy, fx, 360, color=MUT, lw=1.6)
    # index
    arrow(ax, fx + 320, 390, 700, 390, color=CYAN, lw=2.4)
    boxed_label(ax, 700, 300, 200, 80, "bulk_index", "#0E1A30", CYAN, tc=CYAN, sub="hybrid vectors")
    boxed_label(ax, 700, 410, 200, 80, "OpenSearch", BG2, CYAN, sub="enterprise-knowledge")
    arrow(ax, 800, 380, 800, 410, color=CYAN)
    boxed_label(ax, 940, 350, 180, 80, "Agentic RAG", "#0E1A30", GRN, tc=GRN, sub="now grounded\non your wiki")
    arrow(ax, 900, 390, 940, 390, color=GRN)
    ax.text(64, 600, "POST /ingest {\"source\":\"confluence\"}   ·   GET /sources   ·   "
            "python scripts/ingest_sources.py --source sharepoint",
            color=MUT, fontsize=12.5, va="center", fontproperties=MONO)
    footer(ax, 8)
    save(fig, 8)


def s09():
    fig, ax = new_slide()
    title_bar(ax, "Flow", "Connector ingestion — sequence")
    actors = [("CLI / API", BLUE), ("RAG /ingest", CYAN), ("Connector", PUR),
              ("Source system", PINK), ("Vector store", GRN)]
    xs = [140, 380, 620, 860, 1100]
    for (name, c), x in zip(actors, xs):
        boxed_label(ax, x - 90, 180, 180, 50, name, BG2, c, tc=c, title_fs=13)
        ax.plot([x, x], [230, 600], color=LINE, lw=1.4, ls=(0, (4, 4)), zorder=1)
    steps = [
        (0, 1, "ingest(source, options)", CYAN, 270),
        (1, 2, "get_connector() · fetch()", PUR, 320),
        (2, 3, "REST / Graph query  (or sample fallback)", PINK, 370),
        (3, 2, "raw documents (HTML)", MUT, 420),
        (2, 2, "html→text · chunk()", AMB, 470),
        (1, 4, "bulk_index(chunks)", GRN, 520),
        (4, 0, "{documents, chunks, live}", BLUE, 570),
    ]
    for a, b, label, c, y in steps:
        if a == b:
            ax.add_patch(FancyBboxPatch((xs[a] + 10, y - 14), 230, 26,
                         boxstyle="round,pad=0,rounding_size=6", fc="#152038", ec=c, lw=1.4, zorder=3))
            ax.text(xs[a] + 20, y, label, color=c, fontsize=11.5, va="center",
                    zorder=4, fontproperties=MONO)
        else:
            arrow(ax, xs[a], y, xs[b], y, color=c, lw=2)
            mx = (xs[a] + xs[b]) / 2
            ax.text(mx, y - 12, label, color=INK, fontsize=11.5, ha="center",
                    va="center", zorder=4)
    ax.text(64, 630, "Offline-safe: with no credentials the connector returns bundled "
            "sample pages, so the demo always runs.", color=MUT, fontsize=12.5, va="center")
    footer(ax, 9)
    save(fig, 9)


def s10():
    fig, ax = new_slide()
    title_bar(ax, "Governance", "Evaluation gates — quality before release")
    metrics = [("Answer relevancy", "≥ 0.80", GRN), ("Faithfulness", "≥ 0.85", GRN),
               ("Contextual recall", "≥ 0.70", GRN), ("Hallucination", "≤ 0.10", RED),
               ("Toxicity", "≤ 0.05", RED), ("Bias", "≤ 0.10", RED)]
    for i, (m, thr, c) in enumerate(metrics):
        r, col = divmod(i, 2)
        x = 64 + col * 300
        y = 210 + r * 86
        panel(ax, x, y, 280, 70, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((x, y + 10), 5, 50, color=c, zorder=4))
        ax.text(x + 22, y + 30, m, color=INK, fontsize=14, fontweight="bold", va="center")
        ax.text(x + 22, y + 52, "threshold " + thr, color=MUT, fontsize=11.5, va="center")
        chip(ax, x + 200, y + 22, thr, fc="#152038", tc=c, fs=12)
    # judge panel
    panel(ax, 700, 210, 420, 420, fc=BG3, ec=LINE)
    ax.text(722, 244, "LLM-as-judge", color=PUR, fontsize=16, fontweight="bold", va="center")
    bullets(ax, 722, 296, [
        "Weighted rubric across metrics",
        "Self-consistency sampling",
        "Returns weighted_score + pass/fail",
        "Runs per answer (validate stage)",
        "Same suite gates CI deployment",
    ], dy=46, fs=13.5, marker=PUR)
    rbox(ax, 722, 548, 376, 60, fc="#0E2A1E", ec=GRN, lw=2, r=10)
    ax.text(742, 578, "PASS  →  release    ·    FAIL  →  block + log",
            color=GRN, fontsize=14, fontweight="bold", va="center", fontproperties=MONO)
    footer(ax, 10)
    save(fig, 10)


def s11():
    fig, ax = new_slide()
    title_bar(ax, "Modernisation", "AWS Transform Custom — the factory")
    defs = [
        ("enterprise-context-normaliser", "PII redaction · date norm · truncation", CYAN),
        ("java-spring-to-spring-ai", "legacy Spring → Spring AI 1.x", BLUE),
        ("eval-dataset-updater", "regenerate eval cases on schema change", PUR),
        ("fintech-java-spring2-to-spring3", "Boot 2.x → 3.2 · WebClient · OTel", GRN),
        ("fintech-rules-to-agentic-rag", "rule engines → Kiro agentic RAG", ORG),
        ("fintech-soap-to-rest", "SOAP → REST anti-corruption layer", TEAL),
        ("fintech-cve-remediation", "scan → map CVE → patch, gated to 0 HIGH", RED),
    ]
    for i, (name, d, c) in enumerate(defs):
        y = 200 + i * 60
        panel(ax, 64, y, 1056, 50, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((64, y), 6, 50, color=c, zorder=4))
        ax.text(86, y + 25, name, color=INK, fontsize=14, fontweight="bold",
                va="center", fontproperties=MONO)
        ax.text(620, y + 25, d, color=MUT, fontsize=13, va="center")
    ax.text(64, 642, "Driven by AWS Kiro agent specs · 6-service fintech legacy estate "
            "as the modernisation demo", color=CYAN, fontsize=13, va="center")
    footer(ax, 11)
    save(fig, 11)


def s12():
    fig, ax = new_slide()
    title_bar(ax, "Before / after", "Spring AI modernisation")
    # before
    panel(ax, 64, 210, 500, 400, fc="#1A1320", ec=RED)
    ax.add_patch(Rectangle((64, 210), 500, 8, color=RED, zorder=4))
    ax.text(88, 250, "LEGACY  ·  Spring Boot 2.7", color=RED, fontsize=16, fontweight="bold", va="center")
    bullets(ax, 88, 304, [
        "RestTemplate, blocking I/O",
        "Solr search client",
        "Manual prompt strings",
        "No retries / resilience",
        "No telemetry, no eval",
        "CVE-bearing dependencies",
    ], dy=46, fs=14, marker=RED)
    # after
    panel(ax, 616, 210, 504, 400, fc="#0E2A1E", ec=GRN)
    ax.add_patch(Rectangle((616, 210), 504, 8, color=GRN, zorder=4))
    ax.text(640, 250, "MODERN  ·  Spring Boot 3 + Spring AI 1.0", color=GRN,
            fontsize=15, fontweight="bold", va="center")
    bullets(ax, 640, 304, [
        "WebClient, reactive + Resilience4j",
        "AnthropicChatModel (Spring AI)",
        "Grounded prompts from RAG",
        "OTel spans + metrics",
        "Eval gate before response",
        "CVE-remediated, 0 HIGH/CRITICAL",
    ], dy=46, fs=14, marker=GRN)
    arrow(ax, 566, 410, 614, 410, color=CYAN, lw=3)
    footer(ax, 12)
    save(fig, 12)


def s13():
    fig, ax = new_slide()
    title_bar(ax, "Screenshot · mockup", "Single pane of glass — Grafana")
    ix, iy, iw, ih = mock_window(ax, 56, 170, 1168, 470,
                                 url="localhost:3000/d/ai-factory-single-pane")
    # stat row
    stats = [("6", "Services Up", GRN), ("0", "Targets Down", GRN),
             ("12.4/s", "Pipeline req/s", CYAN), ("0.8%", "Error rate", AMB),
             ("0.42s", "Spring p95", BLUE)]
    sw = (iw - 16 * 4) / 5
    for i, (v, l, c) in enumerate(stats):
        stat_tile(ax, ix + i * (sw + 16), iy, sw, 86, v, l, c)
    # two charts
    rng = np.random.default_rng(7)
    def wig(n, base, amp): return np.clip(base + amp * np.cumsum(rng.normal(0, 0.06, n)), 0.05, 0.95)
    mini_chart(ax, ix, iy + 104, iw / 2 - 8, 150,
               [wig(40, 0.5, 1), wig(40, 0.35, 1), wig(40, 0.6, 1)],
               [CYAN, PUR, ORG], "Throughput by stage (req/s)", fill=True)
    mini_chart(ax, ix + iw / 2 + 8, iy + 104, iw / 2 - 8, 150,
               [wig(40, 0.3, 1), wig(40, 0.5, 1)],
               [BLUE, GRN], "p95 span duration (ms)")
    # bottom: error rate chart + table
    mini_chart(ax, ix, iy + 272, iw / 2 - 8, 150,
               [wig(40, 0.12, 0.6), wig(40, 0.06, 0.4)],
               [RED, AMB], "Error rate by service (%)", fill=True)
    panel(ax, ix + iw / 2 + 8, iy + 272, iw / 2 - 8, 150, fc="#0C1424", ec=LINE, r=10, z=5)
    ax.text(ix + iw / 2 + 22, iy + 292, "Service health", color=INK, fontsize=11,
            fontweight="bold", va="center", zorder=7)
    rows = [("spring-service", "UP", GRN), ("rag-service", "UP", GRN),
            ("eval-service", "UP", GRN), ("transform-service", "UP", GRN)]
    for j, (svc, st, c) in enumerate(rows):
        yy = iy + 320 + j * 24
        ax.text(ix + iw / 2 + 22, yy, svc, color=MUT, fontsize=11, va="center",
                zorder=7, fontproperties=MONO)
        chip(ax, ix + iw - 90, yy - 9, st, fc="#0E2A1E", tc=c, fs=9, h=18, padx=8)
    ax.text(56, 668, "Representative mockup — live dashboard needs `docker compose up` "
            "(spanmetrics RED metrics + Spring actuator).", color=MUT, fontsize=11, va="center")
    footer(ax, 13)
    save(fig, 13)


def s14():
    fig, ax = new_slide()
    title_bar(ax, "Telemetry", "Observability pipeline")
    boxed_label(ax, 64, 230, 150, 70, "Spring", BG2, BLUE, sub="actuator")
    boxed_label(ax, 64, 330, 150, 70, "RAG", BG2, CYAN, sub="OTLP spans")
    boxed_label(ax, 64, 430, 150, 70, "Eval / Transform", BG2, GRN, sub="OTLP spans", title_fs=12)
    boxed_label(ax, 320, 330, 200, 90, "OTel Collector", "#0E1A30", CYAN, tc=CYAN,
                sub="spanmetrics\nconnector (RED)")
    for yy in (265, 365, 465):
        arrow(ax, 214, yy, 320, 375, color=MUT, lw=1.8)
    outs = [("Prometheus", AMB, "ai_factory_calls_total\nai_factory_duration_ms"),
            ("Jaeger", PUR, "end-to-end traces"),
            ("Grafana", ORG, "single pane of glass"),
            ("CloudWatch", BLUE, "prod metrics/logs")]
    for i, (t, c, d) in enumerate(outs):
        y = 200 + i * 110
        boxed_label(ax, 620, y, 200, 80, t, BG2, c, tc=c, sub=None, title_fs=15)
        ax.text(840, y + 40, d, color=MUT, fontsize=12, va="center", fontproperties=MONO)
        arrow(ax, 520, 375, 620, y + 40, color=LINE, lw=1.4)
    ax.text(64, 640, "Fix shipped: spanmetrics turns the Python services' traces into "
            "RED metrics so every stage shows up in Grafana.",
            color=CYAN, fontsize=13, va="center")
    footer(ax, 14)
    save(fig, 14)


def s15():
    fig, ax = new_slide()
    title_bar(ax, "Cloud", "AWS deployment architecture")
    panel(ax, 56, 180, 1168, 350, fc=BG3, ec=LINE)
    ax.text(80, 210, "VPC  ·  per environment (dev / sit / preprod / prod)",
            color=ORG, fontsize=14, fontweight="bold", va="center")
    boxed_label(ax, 90, 250, 150, 70, "ALB", BG2, ORG, sub="HTTPS")
    arrow(ax, 240, 285, 300, 285, color=MUT)
    # ECS services
    svcs = [("spring", BLUE), ("rag", CYAN), ("eval", GRN), ("transform", ORG)]
    for i, (t, c) in enumerate(svcs):
        x = 300 + i * 150
        boxed_label(ax, x, 250, 134, 70, t, "#0E1A30", c, tc=c, sub="Fargate +\nADOT", title_fs=13)
    ax.text(300, 350, "ECS Fargate tasks  ·  Cloud Map discovery  ·  ADOT sidecar",
            color=MUT, fontsize=12, va="center")
    # data + supporting
    for i, (t, c) in enumerate([("OpenSearch", CYAN), ("Neo4j + EFS", PUR),
                                ("ECR", ORG), ("Secrets Mgr", AMB), ("CloudWatch", BLUE)]):
        x = 90 + i * 220
        boxed_label(ax, x, 410, 200, 64, t, BG2, c, tc=c, title_fs=13)
    bullets(ax, 64, 580, [
        "One Terraform stack, four env var-files; isolated S3 state per env",
        "Right-sized per environment; canary on preprod/prod",
    ], dy=36, fs=14)
    footer(ax, 15)
    save(fig, 15)


def s16():
    fig, ax = new_slide()
    title_bar(ax, "Delivery", "CI/CD with Harness — gated promotion")
    # CI row
    ax.text(64, 200, "CI  ·  build + scan", color=CYAN, fontsize=14, fontweight="bold", va="center")
    ci = [("pytest + mvn", GRN), ("SAST", AMB), ("Trivy scan", ORG), ("→ ECR", BLUE)]
    x = 64
    for t, c in ci:
        w = 210
        boxed_label(ax, x, 226, w, 58, t, "#0E1A30", c, tc=c, title_fs=13)
        if x > 64:
            arrow(ax, x - 24, 255, x, 255, color=MUT)
        x += w + 24
    # CD row
    ax.text(64, 340, "CD  ·  promote", color=CYAN, fontsize=14, fontweight="bold", va="center")
    cd = [("dev\nrolling", GRN), ("sit\n+ DAST", AMB), ("preprod\ncanary", ORG),
          ("prod\n2-person + canary", RED)]
    x = 64
    for t, c in cd:
        w = 230
        title, sub = t.split("\n")
        boxed_label(ax, x, 366, w, 78, title, BG2, c, tc=c, sub=sub, title_fs=15)
        if x > 64:
            arrow(ax, x - 24, 405, x, 405, color=MUT)
        x += w + 24
    panel(ax, 64, 480, 1056, 150, fc=BG2, ec=LINE)
    ax.text(86, 510, "Promotion gates", color=INK, fontsize=15, fontweight="bold", va="center")
    bullets(ax, 86, 550, [
        "Smoke tests (dev) → ZAP DAST clean + approval (sit)",
        "Canary verified + approval (preprod) → 2-person approval + canary (prod)",
        "Auto-rollback on canary failure (25% → 10-min verify → 100%)",
    ], dy=32, fs=13.5)
    footer(ax, 16)
    save(fig, 16)


def s17():
    fig, ax = new_slide()
    title_bar(ax, "Demo · mockup", "Walkthrough — bring it up & ingest")
    terminal(ax, 56, 175, 700, 470, [
        ("c", "docker compose up -d"),
        ("ok", "✔ 10 services healthy (rag, eval, transform, spring, …)"),
        ("o", ""),
        ("c", "python scripts/seed_data.py --source ./data/sample-docs/"),
        ("ok", "Seeded 42/42 chunks ✓"),
        ("o", ""),
        ("c", "python scripts/ingest_sources.py --list"),
        ("hdr", "Available sources: confluence, sharepoint"),
        ("o", ""),
        ("c", "python scripts/ingest_sources.py --source confluence"),
        ("o", "Ingesting from 'confluence' → http://localhost:8001"),
        ("ok", "Ingested 2 documents → 3 chunks [sample (offline fallback)] ✓"),
        ("o", ""),
        ("c", "python scripts/ingest_sources.py --source sharepoint"),
        ("ok", "Ingested 2 documents → 4 chunks [sample (offline fallback)] ✓"),
    ], title="bash — quick start")
    panel(ax, 784, 175, 440, 470, fc=BG2, ec=LINE)
    ax.text(806, 210, "What just happened", color=INK, fontsize=16, fontweight="bold", va="center")
    bullets(ax, 806, 262, [
        "Full local stack on Docker",
        "Sample corpus seeded into RAG",
        "Confluence + SharePoint connected",
        "Offline fallback → no creds needed",
        "Index now holds wiki + sample docs",
    ], dy=52, fs=13.5, marker=CYAN)
    rbox(ax, 806, 556, 396, 64, fc="#0E2A1E", ec=GRN, lw=2, r=10)
    ax.text(826, 588, "Set CONFLUENCE_* / SHAREPOINT_*\nfor live ingestion",
            color=GRN, fontsize=12.5, va="center", fontproperties=MONO)
    footer(ax, 17)
    save(fig, 17)


def s18():
    fig, ax = new_slide()
    title_bar(ax, "Demo · mockup", "Query → grounded, gated answer")
    terminal(ax, 56, 175, 640, 470, [
        ("c", "curl -X POST localhost:8001/query \\"),
        ("c", "  -d '{\"query\":\"What is our Q2 revenue forecast?\"}'"),
        ("o", ""),
        ("hdr", "{"),
        ("o", '  "answer": "The Q2 APAC revenue forecast'),
        ("o", '            is AUD 4.2M, based on the'),
        ("o", '            confirmed pipeline.",'),
        ("o", '  "sources": [{"title": "Q2 Forecast",'),
        ("o", '     "url": "confluence/.../100021"}],'),
        ("ok", '  "sufficiency_score": 0.91,'),
        ("o", '  "hops": 2,'),
        ("hdr", "}"),
    ], title="bash — query")
    # eval gate card
    panel(ax, 724, 175, 500, 470, fc=BG3, ec=LINE)
    ax.text(748, 212, "Eval gate result", color=GRN, fontsize=16, fontweight="bold", va="center")
    gates = [("Answer relevancy", "0.93", "≥0.80", GRN),
             ("Faithfulness", "0.90", "≥0.85", GRN),
             ("Contextual recall", "0.81", "≥0.70", GRN),
             ("Hallucination", "0.03", "≤0.10", GRN),
             ("Toxicity", "0.00", "≤0.05", GRN)]
    for i, (m, v, thr, c) in enumerate(gates):
        y = 256 + i * 58
        panel(ax, 748, y, 452, 46, fc="#0C1424", ec=LINE, r=8, z=5)
        ax.text(766, y + 23, m, color=INK, fontsize=13, va="center", zorder=7)
        ax.text(1040, y + 23, thr, color=MUT, fontsize=11, va="center", zorder=7,
                fontproperties=MONO)
        chip(ax, 1120, y + 12, v, fc="#0E2A1E", tc=c, fs=11, h=22, padx=8, z=7)
    rbox(ax, 748, 556, 452, 60, fc="#0E2A1E", ec=GRN, lw=2.4, r=10)
    ax.text(774, 586, "✔  PASS — answer released to user",
            color=GRN, fontsize=15, fontweight="bold", va="center")
    footer(ax, 18)
    save(fig, 18)


def s19():
    fig, ax = new_slide()
    title_bar(ax, "Confidence", "Quality, testing & security")
    cards = [("99", "automated tests", "unit + integration, all green", GRN),
             ("0", "ruff lint issues", "consistent, clean code", CYAN),
             ("100%", "mocked I/O", "tests need no infrastructure", PUR),
             ("0", "HIGH/CRITICAL CVEs", "gated by remediation flow", RED)]
    for i, (v, t, d, c) in enumerate(cards):
        x = 64 + i * 270
        panel(ax, x, 200, 246, 150, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((x, 200), 246, 8, color=c, zorder=4))
        ax.text(x + 22, 256, v, color=c, fontsize=34, fontweight="bold", va="center")
        ax.text(x + 22, 300, t, color=INK, fontsize=14, fontweight="bold", va="center")
        ax.text(x + 22, 326, d, color=MUT, fontsize=11.5, va="center")
    panel(ax, 64, 390, 520, 240, fc=BG2, ec=LINE)
    ax.text(86, 422, "Testing strategy", color=INK, fontsize=15, fontweight="bold", va="center")
    bullets(ax, 86, 466, [
        "External services stubbed in conftest",
        "DeepEval suite gated, not run by default",
        "Connector parsing covered by 24 tests",
        "Eval gate doubles as a CI quality gate",
    ], dy=40, fs=13.5, marker=CYAN)
    panel(ax, 604, 390, 516, 240, fc=BG2, ec=LINE)
    ax.text(626, 422, "Security posture", color=INK, fontsize=15, fontweight="bold", va="center")
    bullets(ax, 626, 466, [
        "SAST (Semgrep, Gitleaks, OWASP) in CI",
        "Trivy image scanning before ECR push",
        "OWASP ZAP DAST against SIT",
        "PII redaction in the telemetry pipeline",
    ], dy=40, fs=13.5, marker=GRN)
    footer(ax, 19)
    save(fig, 19)


def s20():
    fig, ax = new_slide()
    for i, (_, c) in enumerate(STAGE):
        ax.add_patch(Rectangle((0, 150 + i * 90), W, 90, color=c, alpha=0.05, zorder=-5))
    grad_band(ax, 0, 8, CYAN)
    title_bar(ax, "Wrap-up", "Roadmap & next steps")
    bullets(ax, 64, 220, [
        "More connectors: Jira, Google Drive, S3, Slack",
        "Graph enrichment on ingest (entities → Neo4j)",
        "Incremental / scheduled sync with change detection",
        "Per-source access control + row-level filtering",
        "Alerting on the single-pane SLOs (error budget burn)",
    ], dy=54, fs=16)
    panel(ax, 700, 200, 420, 300, fc=BG2, ec=LINE)
    ax.text(722, 234, "Delivered this iteration", color=CYAN, fontsize=15,
            fontweight="bold", va="center")
    bullets(ax, 722, 282, [
        "Single pane of glass + RED metrics",
        "Confluence & SharePoint connectors",
        "/ingest API + CLI, 24 new tests",
        "Test-collection fix (99 green)",
    ], dy=46, fs=13.5, marker=GRN)
    ax.text(64, 600, "Next: scaling into a Cognitive Enterprise Platform  →",
            color=CYAN, fontsize=20, fontweight="bold", va="center")
    ax.text(66, 638, "Federated agents · shared memory · A2A orchestration",
            color=MUT, fontsize=13, va="center")
    footer(ax, 20)
    save(fig, 20)


def card_row(ax, x0, x1, y, h, items, color, fs=10, fc="#101A30", lw=1.3, r=8):
    """Evenly spaced mini-cards across [x0, x1]."""
    n = len(items)
    gap = 12
    cw = (x1 - x0 - gap * (n - 1)) / n
    for i, it in enumerate(items):
        cx = x0 + i * (cw + gap)
        rbox(ax, cx, y, cw, h, fc=fc, ec=color, lw=lw, r=r, z=5)
        ax.text(cx + cw / 2, y + h / 2, it, color=INK, fontsize=fs,
                ha="center", va="center", zorder=6)
    return cw


def s21():
    fig, ax = new_slide()
    title_bar(ax, "Target architecture", "The Cognitive Enterprise Platform",
              "Centralised + decentralised agents · long/short-term memory · "
              "A2A orchestration · built bottom-up")

    def lane(y, h, label, color):
        panel(ax, 56, y, 1168, h, fc=BG2, ec=LINE, r=10)
        ax.add_patch(Rectangle((56, y), 6, h, color=color, zorder=4))
        ax.text(74, y + h / 2, label, color=color, fontsize=9.5,
                fontweight="bold", va="center")

    # L1 — business surface
    lane(150, 56, "BUSINESS\nSURFACE", BLUE)
    card_row(ax, 195, 1212, 163, 30,
             ["Finance &\nProcurement", "Operations &\nQuality", "HR & People",
              "Product\nDevelopment", "Supply Chain", "Customer & Sales"],
             BLUE, fs=9)
    # L2 — orchestration
    lane(214, 66, "ORCHESTRA-\nTION", PUR)
    card_row(ax, 195, 1212, 230, 34,
             ["Workflow\nOrchestrator", "Cross-Domain\nCoordinator",
              "Responsible AI\n(RAI)", "Governance\n& Policy",
              "Value & ROI\nMonitor", "Cost & Perf\nEngine"],
             PUR, fs=8.5)
    # L3 — agent ecosystem (two sub-rows)
    lane(288, 160, "AGENT\nECOSYSTEM", CYAN)
    chip(ax, 1054, 296, "A2A protocol", fc="#152038", tc=PINK, fs=9, h=22)
    ax.text(174, 344, "Centrally\nbuilt", color=CYAN, fontsize=9,
            fontweight="bold", va="center")
    card_row(ax, 300, 1212, 322, 44,
             ["RAG &\nKnowledge", "Analytics\nEngine", "Action &\nWorkflow",
              "Visual & Data\nInspector"], CYAN, fs=9, fc="#0E1A30")
    ax.text(174, 406, "Decentrally\nbuilt", color=ORG, fontsize=9,
            fontweight="bold", va="center")
    card_row(ax, 300, 1212, 384, 44,
             ["Invoice &\nAP", "Spend\nAnalyst", "Quality\nControl",
              "Logistics\nCoord.", "Charge\nAssistant", "R&D\nKnowledge"],
             ORG, fs=9, fc="#1F1605")
    # L4 — memory
    lane(456, 82, "MEMORY\nARCH.", GRN)
    ax.text(200, 470, "Long-term", color=GRN, fontsize=9, fontweight="bold",
            va="center")
    card_row(ax, 200, 858, 482, 40,
             ["Org Knowledge\nBase", "Model\nRegistry", "Vector\nDatabase",
              "Institutional\nMemory", "Compliance\nStore"], GRN, fs=8.5)
    ax.text(880, 470, "Short-term", color=AMB, fontsize=9, fontweight="bold",
            va="center")
    card_row(ax, 880, 1212, 482, 40,
             ["Session\nContext", "Scratchpad"], AMB, fs=9)
    # L5 — data foundation & LLMOps
    lane(546, 84, "DATA &\nLLMOPS", ORG)
    card_row(ax, 200, 1212, 568, 40,
             ["Structured\nData", "Unstructured\nData", "Real-time\nStreams",
              "External\nKnowledge", "Embedding\n& Index", "LLMOps\nPipeline"],
             ORG, fs=9)
    ax.text(64, 652, "Build sequence (bottom-up):  data foundation → memory → "
            "agents → orchestration → business value",
            color=CYAN, fontsize=11, fontweight="bold", va="center")
    footer(ax, 21)
    save(fig, 21)


def s22():
    fig, ax = new_slide()
    title_bar(ax, "Business surface", "Value streams across every function",
              "Agents meet people where the work happens — one platform, many domains")
    funcs = [
        ("Finance & Procurement", CYAN,
         ["Invoice & AP automation", "Spend & savings analyst", "Contract Q&A"]),
        ("Operations & Quality", GRN,
         ["Quality control agent", "SOP / runbook copilot", "Incident triage"]),
        ("HR & People", PUR,
         ["Onboarding assistant", "Policy & benefits Q&A", "Talent matching"]),
        ("Product Development", BLUE,
         ["R&D knowledge agent", "Design / code copilot", "Spec drafting"]),
        ("Supply Chain", ORG,
         ["Demand forecasting", "Supplier risk watch", "Logistics coordinator"]),
        ("Customer & Sales", PINK,
         ["Charge / billing assistant", "Proposal generator", "Next-best-action"]),
    ]
    for i, (t, c, items) in enumerate(funcs):
        r, col = divmod(i, 3)
        x = 64 + col * 388
        y = 200 + r * 210
        panel(ax, x, y, 360, 186, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((x, y), 360, 8, color=c, zorder=4))
        ax.text(x + 22, y + 44, t, color=INK, fontsize=15.5,
                fontweight="bold", va="center")
        for j, it in enumerate(items):
            ax.add_patch(Circle((x + 30, y + 92 + j * 34), 4, color=c, zorder=5))
            ax.text(x + 48, y + 92 + j * 34, it, color=MUT, fontsize=12.5, va="center")
    footer(ax, 22)
    save(fig, 22)


def s23():
    fig, ax = new_slide()
    title_bar(ax, "Orchestration layer", "The AI Centre of Excellence",
              "Central control plane that coordinates, governs and optimises every agent")
    caps = [
        ("Workflow orchestration", CYAN,
         "Plan, route and supervise multi-step,\nmulti-agent work end to end"),
        ("Cross-domain coordination", BLUE,
         "Hand off context safely between\nfunction-specific agents"),
        ("Responsible AI (RAI)", PUR,
         "Guardrails, bias & safety checks,\nhuman-in-the-loop policies"),
        ("Governance & policy", GRN,
         "Access control, audit trails,\nmodel & data approvals"),
        ("Value & ROI monitoring", AMB,
         "Track outcomes, adoption and\nbusiness value per agent"),
        ("Cost & performance engine", ORG,
         "Model routing, caching and\ntoken budget guardrails"),
    ]
    for i, (t, c, d) in enumerate(caps):
        r, col = divmod(i, 3)
        x = 64 + col * 388
        y = 210 + r * 200
        panel(ax, x, y, 360, 176, fc=BG2, ec=LINE)
        ax.add_patch(FancyBboxPatch((x + 22, y + 22), 40, 40,
                     boxstyle="round,pad=0,rounding_size=10",
                     fc="#152038", ec=c, lw=2, zorder=4))
        ax.add_patch(Circle((x + 42, y + 42), 8, color=c, zorder=5))
        ax.text(x + 22, y + 92, t, color=INK, fontsize=14.5,
                fontweight="bold", va="center")
        ax.text(x + 22, y + 120, d, color=MUT, fontsize=11.5, va="top")
    footer(ax, 23)
    save(fig, 23)


def s24():
    fig, ax = new_slide()
    title_bar(ax, "Operating model", "Centrally built · decentrally built",
              "A federated model: a paved-road platform that domain teams build on")
    panel(ax, 64, 200, 520, 300, fc="#0E1A30", ec=CYAN)
    ax.add_patch(Rectangle((64, 200), 520, 8, color=CYAN, zorder=4))
    ax.text(88, 242, "Centrally built", color=CYAN, fontsize=18,
            fontweight="bold", va="center")
    ax.text(88, 270, "Platform / CoE team", color=MUT, fontsize=12, va="center")
    bullets(ax, 88, 312, [
        "Shared agents: RAG, analytics, action, inspector",
        "Reusable tools, memory & guardrails",
        "Golden datasets and eval harness",
        "Security, observability, FinOps baked in",
    ], dy=42, fs=13, marker=CYAN)
    panel(ax, 616, 200, 504, 300, fc="#1F1605", ec=ORG)
    ax.add_patch(Rectangle((616, 200), 504, 8, color=ORG, zorder=4))
    ax.text(640, 242, "Decentrally built", color=ORG, fontsize=18,
            fontweight="bold", va="center")
    ax.text(640, 270, "Business / domain teams", color=MUT, fontsize=12, va="center")
    bullets(ax, 640, 312, [
        "Domain agents: invoice, quality, logistics…",
        "Own their prompts, data and use cases",
        "Move fast on the paved road",
        "Compose central agents via A2A",
    ], dy=42, fs=13, marker=ORG)
    panel(ax, 64, 530, 1056, 96, fc=BG2, ec=LINE)
    ax.text(86, 560, "The paved road", color=INK, fontsize=15,
            fontweight="bold", va="center")
    px = 86
    for t, c in [("A2A protocol", PUR), ("Shared memory", GRN),
                 ("Guardrails", RED), ("Eval gates", AMB),
                 ("Observability", CYAN), ("Cost controls", ORG)]:
        px += chip(ax, px, 592, t, fc="#152038", tc=c, fs=12) + 14
    footer(ax, 24)
    save(fig, 24)


def s25():
    fig, ax = new_slide()
    title_bar(ax, "Protocol", "Agent-to-Agent (A2A) collaboration")
    actors = [("Orchestrator", PUR), ("Finance Agent", CYAN), ("RAG Agent", GRN),
              ("Analytics Agent", BLUE), ("Tools / MCP", ORG)]
    xs = [150, 400, 650, 900, 1140]
    for (name, c), x in zip(actors, xs):
        boxed_label(ax, x - 95, 180, 190, 50, name, BG2, c, tc=c, title_fs=13)
        ax.plot([x, x], [230, 590], color=LINE, lw=1.4, ls=(0, (4, 4)), zorder=1)
    steps = [
        (0, 1, "task: 'explain Q2 AP spend spike'", PUR, 272),
        (1, 2, "retrieve(context)", CYAN, 318),
        (2, 4, "hybrid search · graph", GRN, 364),
        (4, 2, "documents + entities", MUT, 410),
        (1, 3, "analyse(trend, anomalies)", CYAN, 456),
        (3, 3, "compute deltas · attribute", BLUE, 502),
        (1, 0, "grounded, attributed answer", GRN, 548),
    ]
    for a, b, label, c, y in steps:
        if a == b:
            ax.add_patch(FancyBboxPatch((xs[a] + 10, y - 14), 240, 26,
                         boxstyle="round,pad=0,rounding_size=6",
                         fc="#152038", ec=c, lw=1.4, zorder=3))
            ax.text(xs[a] + 20, y, label, color=c, fontsize=11.5, va="center",
                    zorder=4, fontproperties=MONO)
        else:
            arrow(ax, xs[a], y, xs[b], y, color=c, lw=2)
            mx = (xs[a] + xs[b]) / 2
            ax.text(mx, y - 12, label, color=INK, fontsize=11, ha="center",
                    va="center", zorder=4)
    ax.text(64, 624, "A2A lets specialised agents delegate sub-tasks with shared "
            "context — orchestrated, observable and policy-checked.",
            color=MUT, fontsize=12.5, va="center")
    footer(ax, 25)
    save(fig, 25)


def s26():
    fig, ax = new_slide()
    title_bar(ax, "Memory", "Memory architecture — what the platform remembers")
    panel(ax, 64, 200, 520, 360, fc=BG2, ec=GRN)
    ax.add_patch(Rectangle((64, 200), 520, 8, color=GRN, zorder=4))
    ax.text(88, 242, "Long-term memory", color=GRN, fontsize=17,
            fontweight="bold", va="center")
    rows = [
        ("Org knowledge base", "docs, wikis, policies → RAG index"),
        ("Vector database", "embeddings for semantic recall"),
        ("Knowledge graph", "entities & relationships (Neo4j)"),
        ("Model & prompt registry", "versioned, governed assets"),
        ("Institutional memory", "past decisions, outcomes, feedback"),
        ("Compliance store", "audit logs, lineage, approvals"),
    ]
    for j, (t, d) in enumerate(rows):
        y = 292 + j * 44
        ax.add_patch(Circle((92, y), 4, color=GRN, zorder=5))
        ax.text(108, y, t, color=INK, fontsize=13, fontweight="bold", va="center")
        ax.text(308, y, d, color=MUT, fontsize=11, va="center")
    panel(ax, 616, 200, 504, 360, fc=BG2, ec=AMB)
    ax.add_patch(Rectangle((616, 200), 504, 8, color=AMB, zorder=4))
    ax.text(640, 242, "Short-term memory", color=AMB, fontsize=17,
            fontweight="bold", va="center")
    rows2 = [
        ("Session context", "current conversation & goals"),
        ("Scratchpad", "intermediate reasoning & plans"),
        ("Tool results cache", "reuse within a single task"),
        ("Working set", "retrieved chunks in play"),
    ]
    for j, (t, d) in enumerate(rows2):
        y = 300 + j * 54
        ax.add_patch(Circle((644, y), 4, color=AMB, zorder=5))
        ax.text(660, y, t, color=INK, fontsize=13.5, fontweight="bold", va="center")
        ax.text(660, y + 22, d, color=MUT, fontsize=11.5, va="center")
    ax.text(64, 600, "Short-term memory is cleared / summarised at task end — and "
            "promoted to long-term when it proves valuable.",
            color=CYAN, fontsize=12.5, va="center")
    footer(ax, 26)
    save(fig, 26)


def s27():
    fig, ax = new_slide()
    title_bar(ax, "Foundation", "Data foundation & LLMOps",
              "From raw enterprise data to a production, observed pipeline")
    srcs = [("Structured data", CYAN, "DBs, ERP, CRM"),
            ("Unstructured data", PUR, "docs, email, wikis"),
            ("Real-time streams", GRN, "events, logs, IoT"),
            ("External knowledge", ORG, "web, market, partners")]
    for i, (t, c, d) in enumerate(srcs):
        boxed_label(ax, 64, 200 + i * 92, 210, 72, t, BG2, c, tc=c, sub=d,
                    title_fs=13)
    stages = [("Ingest &\nclean", AMB), ("Chunk &\nenrich", PUR),
              ("Embed", CYAN), ("Index", GRN), ("Serve", BLUE)]
    x = 340
    for i, (t, c) in enumerate(stages):
        boxed_label(ax, x, 310, 150, 70, t, "#0E1A30", c, tc=c, title_fs=13)
        if i:
            arrow(ax, x - 20, 345, x, 345, color=MUT)
        x += 170
    for i in range(4):
        arrow(ax, 274, 236 + i * 92, 340, 345, color=LINE, lw=1.4)
    panel(ax, 340, 430, 800, 90, fc=BG2, ec=LINE)
    ax.text(362, 460, "LLMOps loop", color=INK, fontsize=14,
            fontweight="bold", va="center")
    px = 362
    for t, c in [("eval gates", GRN), ("drift detection", AMB),
                 ("retraining", PUR), ("versioning", CYAN),
                 ("observability", ORG)]:
        px += chip(ax, px, 490, t, fc="#152038", tc=c, fs=12) + 14
    ax.text(64, 600, "Bottom-up: a solid data + LLMOps base is what makes agents "
            "reliable, current and safe in production.",
            color=CYAN, fontsize=13, va="center")
    footer(ax, 27)
    save(fig, 27)


def s28():
    fig, ax = new_slide()
    title_bar(ax, "Alignment", "How this spike maps to the platform")
    panel(ax, 56, 178, 1168, 56, fc=BG3, ec=LINE)
    ax.text(80, 206, "Platform layer", color=CYAN, fontsize=13,
            fontweight="bold", va="center")
    ax.text(470, 206, "Delivered in this spike", color=GRN, fontsize=13,
            fontweight="bold", va="center")
    ax.text(880, 206, "Next step", color=AMB, fontsize=13,
            fontweight="bold", va="center")
    rows = [
        ("Business surface", "Spring API + query endpoint",
         "Per-function agent UIs", BLUE),
        ("Orchestration", "Kiro 5-stage agent spec",
         "Multi-agent A2A orchestrator", PUR),
        ("Agent ecosystem", "RAG / eval / transform agents",
         "Decentralised domain agents", CYAN),
        ("Memory", "OpenSearch vectors + Neo4j graph",
         "Institutional & session memory", GRN),
        ("Data & LLMOps", "Connectors, seed, eval gates",
         "Streaming + drift + retrain", ORG),
        ("Governance", "Eval gates, PII redaction, SAST/DAST",
         "RAI policies, value/ROI monitor", RED),
    ]
    for i, (l, now, nxt, c) in enumerate(rows):
        y = 246 + i * 62
        panel(ax, 56, y, 1168, 52, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((56, y), 6, 52, color=c, zorder=4))
        ax.text(80, y + 26, l, color=INK, fontsize=13, fontweight="bold",
                va="center")
        ax.text(470, y + 26, now, color=MUT, fontsize=12, va="center")
        ax.text(880, y + 26, nxt, color=MUT, fontsize=12, va="center")
    footer(ax, 28)
    save(fig, 28)


def s29():
    fig, ax = new_slide()
    title_bar(ax, "Trust", "Security, guardrails & Responsible AI")
    cols = [
        ("Input guardrails", CYAN,
         ["Prompt-injection detection", "PII / secret screening",
          "Access & entitlement checks"]),
        ("Output guardrails", GRN,
         ["Groundedness / citation check", "Toxicity, bias, safety filters",
          "Eval gate before release"]),
        ("Platform controls", PUR,
         ["RBAC + audit trails", "Data lineage & approvals",
          "Human-in-the-loop for high risk"]),
    ]
    for i, (t, c, items) in enumerate(cols):
        x = 64 + i * 388
        panel(ax, x, 200, 360, 300, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((x, 200), 360, 8, color=c, zorder=4))
        ax.text(x + 24, 244, t, color=INK, fontsize=16,
                fontweight="bold", va="center")
        for j, it in enumerate(items):
            ax.add_patch(Circle((x + 32, 300 + j * 52), 4, color=c, zorder=5))
            ax.text(x + 50, 300 + j * 52, it, color=MUT, fontsize=12.5, va="center")
    panel(ax, 64, 530, 1056, 96, fc=BG2, ec=LINE)
    ax.text(86, 560, "Defence in depth", color=INK, fontsize=15,
            fontweight="bold", va="center")
    ax.text(86, 592, "Every request passes input → policy → retrieval → "
            "generation → output gates, with full telemetry and audit at each hop.",
            color=MUT, fontsize=12.5, va="center")
    footer(ax, 29)
    save(fig, 29)


def s30():
    fig, ax = new_slide()
    title_bar(ax, "FinOps", "Cost & performance engine",
              "Make agents fast and affordable at scale")
    tiles = [("-62%", "cost per query", "via caching + routing", GRN),
             ("38%", "cache hit rate", "semantic + exact", CYAN),
             ("3", "model tiers", "route by difficulty", PUR),
             ("100%", "budget-guarded", "hard token caps", AMB)]
    for i, (v, t, d, c) in enumerate(tiles):
        x = 64 + i * 270
        panel(ax, x, 200, 246, 150, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((x, 200), 246, 8, color=c, zorder=4))
        ax.text(x + 22, 256, v, color=c, fontsize=32, fontweight="bold", va="center")
        ax.text(x + 22, 300, t, color=INK, fontsize=14, fontweight="bold", va="center")
        ax.text(x + 22, 326, d, color=MUT, fontsize=11, va="center")
    panel(ax, 64, 390, 520, 240, fc=BG2, ec=LINE)
    ax.text(86, 422, "Levers", color=INK, fontsize=15, fontweight="bold", va="center")
    bullets(ax, 86, 466, [
        "Route easy calls to small models",
        "Cache embeddings & frequent answers",
        "Batch + stream to cut latency",
        "Trim context with normalisation",
    ], dy=40, fs=13, marker=CYAN)
    panel(ax, 604, 390, 516, 240, fc=BG2, ec=LINE)
    ax.text(626, 422, "Governance", color=INK, fontsize=15,
            fontweight="bold", va="center")
    bullets(ax, 626, 466, [
        "Per-team token budgets & alerts",
        "Cost attributed by agent & domain",
        "ROI tracked against business value",
        "Anomaly detection on spend",
    ], dy=40, fs=13, marker=AMB)
    footer(ax, 30)
    save(fig, 30)


def s31():
    fig, ax = new_slide()
    title_bar(ax, "Adoption", "Personas & journeys")
    ps = [
        ("Knowledge worker", CYAN, "Asks a question",
         "Gets a grounded, cited answer in seconds"),
        ("Domain builder", ORG, "Builds an agent",
         "Composes platform agents on the paved road"),
        ("Platform engineer", PUR, "Operates the platform",
         "Ships shared agents, memory & guardrails"),
        ("Risk & compliance", GRN, "Governs usage",
         "Audits, approves and monitors every agent"),
    ]
    for i, (t, c, a, b) in enumerate(ps):
        r, col = divmod(i, 2)
        x = 64 + col * 548
        y = 210 + r * 210
        panel(ax, x, y, 520, 186, fc=BG2, ec=LINE)
        ax.add_patch(FancyBboxPatch((x + 24, y + 24), 48, 48,
                     boxstyle="round,pad=0,rounding_size=12",
                     fc="#152038", ec=c, lw=2, zorder=4))
        ax.add_patch(Circle((x + 48, y + 48), 10, color=c, zorder=5))
        ax.text(x + 92, y + 40, t, color=INK, fontsize=17,
                fontweight="bold", va="center")
        ax.text(x + 92, y + 72, a, color=c, fontsize=12.5, va="center")
        ax.text(x + 24, y + 124, b, color=MUT, fontsize=12.5, va="center")
    footer(ax, 31)
    save(fig, 31)


def s32():
    fig, ax = new_slide()
    title_bar(ax, "Maturity", "From pilot to platform")
    stages = [
        ("Crawl", CYAN, "Single agent, single domain", "RAG + eval on one use case"),
        ("Walk", BLUE, "Several agents, shared memory", "Connectors, observ., gates"),
        ("Run", PUR, "Federated agent ecosystem", "A2A across functions"),
        ("Fly", GRN, "Self-optimising platform", "Auto-eval, auto-route, auto-scale"),
    ]
    for i, (t, c, a, b) in enumerate(stages):
        x = 64 + i * 290
        y = 470 - i * 84
        panel(ax, x, y, 260, 150, fc=BG2, ec=c)
        ax.add_patch(Rectangle((x, y), 260, 8, color=c, zorder=4))
        ax.text(x + 22, y + 44, t, color=c, fontsize=20, fontweight="bold", va="center")
        ax.text(x + 22, y + 84, a, color=INK, fontsize=12, fontweight="bold", va="center")
        ax.text(x + 22, y + 112, b, color=MUT, fontsize=11, va="top")
        if i:
            arrow(ax, x - 26, y + 150, x + 8, y + 78, color=MUT, lw=2)
    ax.text(64, 650, "We are here:  Walk — connectors, single pane of glass and "
            "eval gates are live.", color=AMB, fontsize=13,
            fontweight="bold", va="center")
    footer(ax, 32)
    save(fig, 32)


def s33():
    fig, ax = new_slide()
    title_bar(ax, "Plan", "Implementation roadmap")
    cols = [
        ("Now · 0-3 mo", GRN,
         ["Harden RAG + connectors", "Single pane of glass SLOs",
          "2-3 lighthouse use cases", "Eval gates in CI/CD"]),
        ("Next · 3-9 mo", AMB,
         ["A2A multi-agent orchestrator", "Institutional + session memory",
          "Domain agent toolkit", "FinOps cost engine"]),
        ("Later · 9-18 mo", PUR,
         ["Federated agent marketplace", "Streaming + drift + retrain",
          "Self-optimising routing", "Org-wide rollout"]),
    ]
    for i, (t, c, items) in enumerate(cols):
        x = 64 + i * 388
        panel(ax, x, 200, 360, 420, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((x, 200), 360, 8, color=c, zorder=4))
        ax.text(x + 24, 246, t, color=c, fontsize=17, fontweight="bold", va="center")
        bullets(ax, x + 24, 300, items, dy=54, fs=13.5, marker=c)
    footer(ax, 33)
    save(fig, 33)


def s34():
    fig, ax = new_slide()
    title_bar(ax, "Value", "Business value & ROI")
    tiles = [("3-5x", CYAN, "faster knowledge retrieval"),
             ("40-60%", GRN, "less manual rework"),
             ("weeks→days", PUR, "legacy modernisation"),
             ("24/7", ORG, "consistent, audited answers")]
    for i, (v, c, d) in enumerate(tiles):
        x = 64 + i * 270
        panel(ax, x, 200, 246, 140, fc=BG2, ec=LINE)
        ax.add_patch(Rectangle((x, 200), 246, 8, color=c, zorder=4))
        ax.text(x + 22, 256, v, color=c, fontsize=25, fontweight="bold", va="center")
        ax.text(x + 22, 300, d, color=MUT, fontsize=12, va="top")
    panel(ax, 64, 380, 1056, 240, fc=BG2, ec=LINE)
    ax.text(86, 414, "Value levers", color=INK, fontsize=16,
            fontweight="bold", va="center")
    levers = [("Productivity", CYAN, "agents handle the busywork"),
              ("Quality", GRN, "grounded, gated, consistent"),
              ("Speed", PUR, "modernise & ship faster"),
              ("Risk", RED, "governed, auditable, safe"),
              ("Cost", ORG, "routing + caching + budgets"),
              ("Scale", BLUE, "one platform, many domains")]
    for i, (t, c, d) in enumerate(levers):
        r, col = divmod(i, 3)
        x = 86 + col * 350
        y = 462 + r * 74
        ax.add_patch(Circle((x + 8, y), 6, color=c, zorder=5))
        ax.text(x + 28, y, t, color=INK, fontsize=14, fontweight="bold", va="center")
        ax.text(x + 28, y + 24, d, color=MUT, fontsize=11.5, va="center")
    footer(ax, 34)
    save(fig, 34)


def s35():
    fig, ax = new_slide()
    for i, (_, c) in enumerate(STAGE):
        ax.add_patch(Rectangle((0, 150 + i * 90), W, 90, color=c, alpha=0.05, zorder=-5))
    grad_band(ax, 0, 8, CYAN)
    ax.text(64, 250, "From spike to", color=INK, fontsize=46,
            fontweight="bold", va="center")
    ax.text(64, 316, "Cognitive Enterprise Platform", color=CYAN, fontsize=46,
            fontweight="bold", va="center")
    ax.text(66, 372, "Grounded · governed · observable · cost-aware AI — "
            "built bottom-up, scaled org-wide.", color=MUT, fontsize=16, va="center")
    px = 66
    for t, c in [("Agentic RAG", CYAN), ("A2A orchestration", PUR),
                 ("Shared memory", GRN), ("Eval gates", AMB), ("FinOps", ORG)]:
        px += chip(ax, px, 420, t, fc="#152038", tc=c, fs=12) + 12
    panel(ax, 64, 470, 1056, 110, fc=BG2, ec=LINE)
    ax.text(86, 502, "Call to action", color=INK, fontsize=15,
            fontweight="bold", va="center")
    ax.text(86, 536, "Pick 2-3 lighthouse use cases · stand up the paved road · "
            "prove ROI · then federate across the business.",
            color=MUT, fontsize=13, va="center")
    ax.text(64, 624, "Thank you  ·  Questions & live demo", color=INK,
            fontsize=22, fontweight="bold", va="center")
    footer(ax, 35)
    save(fig, 35)


def main():
    for fn in [s01, s02, s03, s04, s05, s06, s07, s08, s09, s10,
               s11, s12, s13, s14, s15, s16, s17, s18, s19, s20,
               s21, s22, s23, s24, s25, s26, s27, s28, s29, s30,
               s31, s32, s33, s34, s35]:
        fn()
    # Assemble PDF
    from PIL import Image
    imgs = [Image.open(p).convert("RGB") for p in _slides]
    pdf = HERE / "ai-factory-demo.pdf"
    imgs[0].save(pdf, save_all=True, append_images=imgs[1:])
    print(f"PDF  → {pdf}  ({len(imgs)} slides)")
    # Assemble PPTX (full-bleed image per slide)
    try:
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]
        for p in _slides:
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(p), 0, 0,
                                     width=prs.slide_width, height=prs.slide_height)
        pptx = HERE / "ai-factory-demo.pptx"
        prs.save(pptx)
        print(f"PPTX → {pptx}")
    except Exception as e:                                  # pragma: no cover
        print("PPTX skipped:", e)


if __name__ == "__main__":
    main()
